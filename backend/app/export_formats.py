"""
Disa aktarim bicimleri: CSV, Excel (.xlsx), PDF.

Uc bicim ayni satirlari gosteriyor (build_rows): kullanici Excel'de
gordugu telefonu PDF'te de bulmali. Yalnizca sunum farkli:
- CSV: baska sistemlere aktarim icin duz veri.
- Excel: gunluk calisma; sabit baslik, filtre, tiklanabilir baglantilar.
- PDF: paylasma/yazdirma; tablo gorunumu, sayfa numarasi, kaynak atfi.
"""

import csv
import io
import re
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

# OSM iletisim bilgisini iki semayla tutuyor: duz (`phone`) ve `contact:`
# onekli (`contact:phone`). Sira oncelik demek. Ayni oncelik listesi
# arayuzde src/lib/contact.ts icinde yasiyor; ikisi ayni sonucu vermeli,
# yoksa kullanici ekranda gordugu numarayi dosyada bulamaz.
PHONE_KEYS = ("phone", "contact:phone", "telephone", "contact:mobile", "mobile")
EMAIL_KEYS = ("email", "contact:email")
WEBSITE_KEYS = ("website", "contact:website", "url", "contact:url")

# Excel, Turkce yerelde liste ayiracini ';' olarak okuyor; virgulle
# ayrilmis dosya tek sutuna yigiliyor.
CSV_DELIMITER = ";"

# Arayuzdeki src/lib/labels.ts ile ayni; kullanici teknik olmayan gonullu,
# dosyada "kindergarten" degil "Anaokulu" gormeli.
TYPE_LABELS = {
    "factory": "Fabrika",
    "office": "Ofis",
    "workshop": "Atölye",
    "kindergarten": "Anaokulu",
    "primary_school": "İlkokul",
    "middle_school": "Ortaokul",
    "high_school": "Lise",
    "private_school": "Özel Okul",
    "college_keyword": "Kolej",
    "college_university": "Üniversite",
    "hotel": "Otel",
    "company": "Şirket",
    "holding": "Holding",
    "real_estate": "Emlak Ofisi",
    "language_school": "Dil Kursu",
    "travel_agency": "Seyahat Acentesi",
    "zoo_aquarium": "Hayvanat Bahçesi & Akvaryum",
    "theme_park": "Tema & Su Parkı",
    "museum": "Müze",
    "botanical_garden": "Botanik Bahçesi",
    "nature_park": "Milli Park & Doğa Alanı",
}

# Arayuzde bos adres bu metinle gosteriliyor; dosyaya tasinmamali.
EMPTY_ADDRESS = "Adres bilgisi yok"

CSV_HEADER = [
    "Ad",
    "Tür",
    "Telefon",
    "E-posta",
    "Web Sitesi",
    "Adres",
    "Konum",
    "Harita",
    "Kayıt No",
]

# OSM verisi ODbL lisansli: paylasilan her kopyada atif zorunlu.
ATTRIBUTION = "Veri: © OpenStreetMap katkıcıları (ODbL) · Overture Maps Foundation"


def pick_tag(tags: Dict[str, Any], keys: tuple) -> str:
    """Return the first non-empty tag value following the given priority."""
    for key in keys:
        value = (tags.get(key) or "").strip()
        if value:
            return value
    return ""


def format_phone(raw: str) -> str:
    """
    Telefonu okunur ve Excel'in sayi sanmayacagi bicime getirir.

    "+905424703486" Excel'de 9,05E+11 oluyor. Bosluklu yazim hem gozle
    okunur hem de metin olarak kalir. Turk numaralari (+90 / 0 onekli,
    10 hane) "+90 542 470 34 86" seklinde; taninmayanlar oldugu gibi.
    """
    raw = raw.strip()
    if not raw:
        return ""
    digits = re.sub(r"\D", "", raw)
    if digits.startswith("90") and len(digits) == 12:
        digits = digits[2:]
    elif digits.startswith("0") and len(digits) == 11:
        digits = digits[1:]
    if len(digits) == 10:
        return f"+90 {digits[:3]} {digits[3:6]} {digits[6:8]} {digits[8:]}"
    return raw


@dataclass(frozen=True)
class ExportRow:
    name: str
    type_label: str
    phone: str
    email: str
    website: str
    address: str
    # Tek hucre metin: Turkce Excel "28.86" degerini binlik ayirac sanip
    # bozuyor; "41.030743, 28.860546" ise oldugu gibi kalir.
    location: str
    maps_url: str
    record_id: str

    def as_list(self) -> List[str]:
        return [
            self.name,
            self.type_label,
            self.phone,
            self.email,
            self.website,
            self.address,
            self.location,
            self.maps_url,
            self.record_id,
        ]


def build_rows(items: List[Dict[str, Any]]) -> List[ExportRow]:
    """Istemciden gelen yerleri uc bicimin ortak satirlarina cevirir."""
    rows = []
    for item in items:
        tags = item.get("tags", {}) or {}
        lat, lon = item.get("lat"), item.get("lon")
        has_coords = isinstance(lat, (int, float)) and isinstance(lon, (int, float))
        address = (item.get("address") or tags.get("addr:full") or "").strip()
        place_type = item.get("type") or ""
        rows.append(
            ExportRow(
                name=item.get("name") or "",
                type_label=TYPE_LABELS.get(place_type, place_type),
                phone=format_phone(pick_tag(tags, PHONE_KEYS)),
                email=pick_tag(tags, EMAIL_KEYS),
                website=pick_tag(tags, WEBSITE_KEYS),
                address="" if address == EMPTY_ADDRESS else address,
                location=f"{lat:.6f}, {lon:.6f}" if has_coords else "",
                maps_url=f"https://www.google.com/maps?q={lat:.6f},{lon:.6f}" if has_coords else "",
                record_id=item.get("id") or "",
            )
        )
    return rows


# --- CSV --------------------------------------------------------------------


def build_csv(items: List[Dict[str, Any]]) -> str:
    """Construct CSV string from OSM items."""
    output = io.StringIO()
    writer = csv.writer(output, delimiter=CSV_DELIMITER)
    writer.writerow(CSV_HEADER)
    for row in build_rows(items):
        writer.writerow(row.as_list())
    return output.getvalue()


def to_csv(items: List[Dict[str, Any]]) -> bytes:
    # BOM sart: Excel BOM'suz bir CSV'yi Windows'ta sistem kod sayfasiyla
    # aciyor ve Turkce karakterler bozuluyor ("İstanbul" -> "Ä°stanbul").
    return ("\ufeff" + build_csv(items)).encode("utf-8")


# --- Ortak ------------------------------------------------------------------

# Marka renkleri (src/app/tokens.css, OKLCH -> sRGB).
INK = "#192029"
INK_2 = "#323841"
INK_3 = "#656b73"
INK_4 = "#81878d"
RULE = "#dbdee2"
PAPER_2 = "#f1f4f7"
PAPER_3 = "#e8ebef"
ACCENT = "#0064da"


def _web_label(url: str) -> str:
    """
    https://www.ornek.com.tr/odalar?x=1 -> ornek.com.tr

    Tabloda yalnizca alan adi: tam adres (Hilton'un sorgu parametreli
    adresi) hucrede 7 satira yayiliyordu. Baglanti yine tam adrese gider.
    """
    host = re.sub(r"^(https?://)?(www\.)?", "", url.strip())
    return re.split(r"[/?#]", host, maxsplit=1)[0]


def _web_href(url: str) -> str:
    return url if re.match(r"^https?://", url) else f"https://{url}"


def _summary(rows: List[ExportRow]) -> str:
    """'Otel 136 · Emlak Ofisi 99 · ...' - tek turse bos."""
    counts = Counter(r.type_label for r in rows if r.type_label)
    if len(counts) < 2:
        return ""
    return " · ".join(f"{label} {n}" for label, n in counts.most_common())


def _turkish_date(when: datetime) -> str:
    months = [
        "Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran",
        "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık",
    ]
    return f"{when.day} {months[when.month - 1]} {when.year}"


# --- Excel ------------------------------------------------------------------


def to_xlsx(items: List[Dict[str, Any]], title: str, generated_at: datetime) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = build_rows(items)
    wb = Workbook()
    ws = wb.active
    ws.title = "Kayıtlar"

    ws.append(CSV_HEADER)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=INK.lstrip("#"))
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    link_font = Font(color=ACCENT.lstrip("#"), underline="single")
    values = []
    # Satir numarasi sayactan: ws.max_row her cagrida tum hucreleri tariyor
    # ve dongude O(n^2) oluyordu (10.000 satir 51 sn).
    for r, row in enumerate(rows, start=2):
        line = [
            row.name,
            row.type_label,
            row.phone,
            row.email,
            row.website,
            row.address,
            row.location,
            "Haritada aç" if row.maps_url else "",
            row.record_id,
        ]
        values.append(line)
        ws.append(line)
        links = {
            3: f"tel:{row.phone.replace(' ', '')}" if row.phone else "",
            4: f"mailto:{row.email}" if row.email else "",
            5: _web_href(row.website) if row.website else "",
            8: row.maps_url,
        }
        for col, href in links.items():
            if href:
                cell = ws.cell(row=r, column=col)
                cell.hyperlink = href
                cell.font = link_font

    # Baslik kaydirirken gorunur kalsin; her sutunda filtre hazir.
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CSV_HEADER))}{len(rows) + 1}"

    # Genislik icerige gore, ama adres gibi uzun sutunlar ekrani kaplamasin.
    for idx, header in enumerate(CSV_HEADER):
        width = max([len(header)] + [len(line[idx] or "") for line in values]) + 2
        ws.column_dimensions[get_column_letter(idx + 1)].width = min(max(width, 8), 50)

    info = wb.create_sheet("Bilgi")
    for line in (
        [title],
        [f"{len(rows)} kayıt · {_turkish_date(generated_at)}"],
        [_summary(rows)],
        [],
        [ATTRIBUTION],
    ):
        info.append(line)
    info["A1"].font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 90

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --- PDF --------------------------------------------------------------------

FONT_DIR = Path(__file__).parent / "fonts"
_fonts_registered = False


def _register_fonts() -> None:
    """Marka yazi tipleri (Inter, Space Grotesk; OFL). Turkce harfler icin
    gomulu TTF sart: PDF'in standart Helvetica'sinda s, g, i yok."""
    global _fonts_registered
    if _fonts_registered:
        return
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    pdfmetrics.registerFont(TTFont("Inter", FONT_DIR / "Inter-Regular.ttf"))
    pdfmetrics.registerFont(TTFont("Inter-SemiBold", FONT_DIR / "Inter-SemiBold.ttf"))
    pdfmetrics.registerFont(TTFont("SpaceGrotesk", FONT_DIR / "SpaceGrotesk-SemiBold.ttf"))
    addMapping("Inter", 0, 0, "Inter")
    addMapping("Inter", 1, 0, "Inter-SemiBold")
    _fonts_registered = True


# Tablo sutunlari: (anahtar, baslik, genislik pt). Toplam = yatay A4 -
# kenar boslugu (770). Konum ve kayit no PDF'te yok: ad hucresi zaten
# haritaya baglantili.
PDF_COLUMNS = [
    ("num", "#", 26),
    ("name", "AD", 176),
    ("type", "TÜR", 82),
    ("phone", "TELEFON", 92),
    ("email", "E-POSTA", 128),
    ("web", "WEB SİTESİ", 108),
    ("address", "ADRES", 158),
]
# Tek turlu listede TUR sutunu tekrar ("Otel, Otel, ..."); tur basliktan
# belli. Acilan yer en cok tasan iki sutuna gidiyor.
PDF_COLUMNS_SINGLE_TYPE = [
    (key, label, {"email": 168, "address": 200}.get(key, width))
    for key, label, width in PDF_COLUMNS
    if key != "type"
]


# Bos hucre isareti ve duz metin hucrelerin satir yuksekligi (Paragraph
# stilleriyle ayni, 7.6 pt yazi).
EMPTY = "—"
STRING_LEADING = 9.6


def pdf_columns(rows: List[ExportRow]) -> list:
    """Tek turlu listede TUR sutunu yok (bkz. PDF_COLUMNS_SINGLE_TYPE)."""
    return PDF_COLUMNS_SINGLE_TYPE if len({r.type_label for r in rows}) <= 1 else PDF_COLUMNS


def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def to_pdf(
    items: List[Dict[str, Any]],
    title: str,
    generated_at: datetime,
    subtitle: Optional[str] = None,
) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfgen import canvas as pdf_canvas
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    _register_fonts()
    rows = build_rows(items)
    page_w, page_h = landscape(A4)
    margin_x, margin_top, margin_bottom = 34, 36, 34

    def style(name, **kw) -> ParagraphStyle:
        base = dict(fontName="Inter", fontSize=7.6, leading=9.6, textColor=colors.HexColor(INK))
        base.update(kw)
        return ParagraphStyle(name, **base)

    eyebrow = style("eyebrow", fontName="Inter-SemiBold", fontSize=7, textColor=colors.HexColor(INK_3))
    h1 = style("h1", fontName="SpaceGrotesk", fontSize=19, leading=23)
    meta = style("meta", fontSize=8.5, leading=11, textColor=colors.HexColor(INK_2))
    summary = style("summary", fontSize=7.5, leading=10, textColor=colors.HexColor(INK_3))
    cell = style("cell")
    cell_name = style("cell_name", fontName="Inter-SemiBold")
    head = style("head", fontName="Inter-SemiBold", fontSize=6.6, textColor=colors.HexColor(INK_3))

    def linked(text: str, href: str, st=cell, color=ACCENT) -> Paragraph:
        return Paragraph(f'<a href="{_esc(href)}" color="{color}">{_esc(text)}</a>', st)

    with_phone = sum(1 for r in rows if r.phone)
    ratio = f" (%{round(100 * with_phone / len(rows))})" if rows else ""
    story = [
        Paragraph("ROTA · KURUM LİSTESİ", eyebrow),
        Spacer(0, 3),
        Paragraph(_esc(title), h1),
        Spacer(0, 3),
        Paragraph(
            _esc(
                f"{len(rows)} kayıt · telefonu olan {with_phone}{ratio} · "
                f"{_turkish_date(generated_at)}" + (f" · {subtitle}" if subtitle else "")
            ),
            meta,
        ),
    ]
    type_summary = _summary(rows)
    if type_summary:
        story += [Spacer(0, 2), Paragraph(_esc(type_summary), summary)]
    story.append(Spacer(0, 12))

    columns = pdf_columns(rows)
    col_widths = [w for _, _, w in columns]

    # Bagli olmayan kisa hucreler (sira no, bos "—") duz metin: Paragraph
    # olcumu satir basina en pahali is.
    def cells(i: int, r: ExportRow) -> dict:
        return {
            "num": str(i),
            "name": linked(r.name or "İsimsiz yer", r.maps_url, cell_name, INK)
            if r.maps_url
            else Paragraph(_esc(r.name or "İsimsiz yer"), cell_name),
            "type": Paragraph(_esc(r.type_label), cell) if r.type_label else EMPTY,
            "phone": linked(r.phone, f"tel:{r.phone.replace(' ', '')}") if r.phone else EMPTY,
            "email": linked(r.email, f"mailto:{r.email}") if r.email else EMPTY,
            "web": linked(_web_label(r.website), _web_href(r.website)) if r.website else EMPTY,
            "address": Paragraph(_esc(r.address), cell) if r.address else EMPTY,
        }

    header = [Paragraph(label, head) for _, label, _ in columns]
    body = []
    for i, r in enumerate(rows, start=1):
        row_cells = cells(i, r)
        body.append([row_cells[key] for key, _, _ in columns])

    # Sayfalama elle: tek bir LongTable'i kutuphane bolerken her sayfa
    # sonunda kalan TUM satirlarin yuksekligini yeniden hesapliyordu
    # (4000 satir 28 sn, karesel). Burada her satir bir kez olculuyor,
    # sayfalara doldurulup her sayfaya kendi basligiyla ayri tablo konuyor.
    pad_y, pad_x = 4.5, 5

    def row_height(row) -> float:
        tallest = STRING_LEADING
        for value, width in zip(row, col_widths):
            if not isinstance(value, str):
                tallest = max(tallest, value.wrap(width - 2 * pad_x, 1e6)[1])
        return tallest + 2 * pad_y

    frame_w = page_w - 2 * margin_x - 12  # SimpleDocTemplate cercevesi 6 pt ic bosluk
    frame_h = page_h - margin_top - margin_bottom - 12
    intro_h = sum(f.wrap(frame_w, frame_h)[1] for f in story)
    header_h = row_height(header)
    safety = 4  # olcum ile tablo yerlesimi arasindaki yuvarlama payi

    pages: list[list] = []
    current: list = []
    used = intro_h + header_h
    for row in body:
        h = row_height(row)
        if current and used + h > frame_h - safety:
            pages.append(current)
            current, used = [], header_h
        current.append(row)
        used += h
    if current or not pages:
        pages.append(current)

    zebra = [colors.white, colors.HexColor(PAPER_2)]
    start = 0
    for index, page_rows in enumerate(pages):
        muted = [
            ("TEXTCOLOR", (c, r), (c, r), colors.HexColor(INK_4))
            for r, row in enumerate(page_rows, start=1)
            for c, value in enumerate(row)
            if isinstance(value, str) and value == EMPTY
        ]
        table = Table([header] + page_rows, colWidths=col_widths, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), pad_y),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), pad_y),
                    ("LEFTPADDING", (0, 0), (-1, -1), pad_x),
                    ("RIGHTPADDING", (0, 0), (-1, -1), pad_x),
                    ("FONTNAME", (0, 0), (-1, -1), "Inter"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7.6),
                    ("LEADING", (0, 0), (-1, -1), STRING_LEADING),
                    ("TEXTCOLOR", (0, 1), (0, -1), colors.HexColor(INK_4)),
                    ("ALIGN", (0, 1), (0, -1), "RIGHT"),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(PAPER_3)),
                    ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor(INK_4)),
                    # Zebra sayfalar arasinda kesintisiz: tek numarali
                    # satirla baslayan sayfada renk sirasi donuyor.
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), zebra[start % 2 :] + zebra[: start % 2]),
                    ("LINEBELOW", (0, 1), (-1, -1), 0.3, colors.HexColor(RULE)),
                    *muted,
                ]
            )
        )
        story.append(table)
        if index < len(pages) - 1:
            story.append(PageBreak())
        start += len(page_rows)

    class NumberedCanvas(pdf_canvas.Canvas):
        """'Sayfa x / y' icin toplam sayfa gerekli: once hepsi toplanir."""

        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._pages = []

        def showPage(self):  # noqa: N802 - reportlab API'sini eziyor
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._pages)
            for state in self._pages:
                self.__dict__.update(state)
                self._decorate(total)
                super().showPage()
            super().save()

        def _decorate(self, total: int) -> None:
            y = margin_bottom - 16
            self.setStrokeColor(colors.HexColor(RULE))
            self.setLineWidth(0.4)
            self.line(margin_x, y + 9, page_w - margin_x, y + 9)
            self.setFont("Inter", 6.5)
            self.setFillColor(colors.HexColor(INK_4))
            self.drawString(margin_x, y, ATTRIBUTION)
            self.drawRightString(page_w - margin_x, y, f"Sayfa {self._pageNumber} / {total}")
            if self._pageNumber > 1:
                # Sonraki sayfalarda hangi listeye bakildigi kaybolmasin.
                self.setFont("Inter-SemiBold", 7)
                self.setFillColor(colors.HexColor(INK_3))
                self.drawString(margin_x, page_h - margin_top + 14, title)

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=(page_w, page_h),
        leftMargin=margin_x,
        rightMargin=margin_x,
        topMargin=margin_top,
        bottomMargin=margin_bottom,
        title=title,
        author="Rota",
        subject="Ortak kurum listesi",
    )
    doc.build(story, canvasmaker=NumberedCanvas)
    return out.getvalue()
