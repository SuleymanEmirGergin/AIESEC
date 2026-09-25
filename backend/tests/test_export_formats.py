"""
Disa aktarim bicimleri: CSV, Excel, PDF.

Uc bicim ayni satirlardan uretiliyor; testler her bicimin kendi
vaadine bakiyor: Excel'de sabit baslik + filtre + tiklanabilir baglanti,
PDF'te gecerli dosya + gomulu Turkce yazi tipi + tek turde TUR sutunu yok.
"""

import io
from datetime import datetime, timezone

from openpyxl import load_workbook

from app.export_formats import (
    CSV_HEADER,
    _web_label,
    build_rows,
    pdf_columns,
    to_csv,
    to_pdf,
    to_xlsx,
)

NOW = datetime(2026, 9, 25, tzinfo=timezone.utc)


def _item(i: int, place_type: str = "hotel", **tags) -> dict:
    return {
        "id": f"osm:node:{i}",
        "name": f"Şişli Öğretmen Evi {i}",
        "type": place_type,
        "lat": 41.0 + i / 1000,
        "lon": 29.0,
        "address": "Bağdat Cad. No:1",
        "tags": tags,
    }


ITEMS = [
    _item(1, phone="+902163334455", email="info@ornek.com", website="https://www.ornek.com.tr/odalar?x=1"),
    _item(2, place_type="high_school"),
]


class TestOrtakSatirlar:
    def test_uc_bicim_ayni_satiri_goruyor(self):
        row = build_rows(ITEMS)[0]
        assert row.phone == "+90 216 333 44 55"
        assert row.type_label == "Otel"
        assert row.maps_url.startswith("https://www.google.com/maps?q=41.001000,29.000000")

    def test_web_etiketi_yalnizca_alan_adi(self):
        """Tam adres PDF hucresinde 7 satira yayiliyordu."""
        assert _web_label("https://www.ornek.com.tr/odalar?x=1") == "ornek.com.tr"
        assert _web_label("ornek.com") == "ornek.com"


class TestCsv:
    def test_bom_ve_noktali_virgul(self):
        data = to_csv(ITEMS).decode("utf-8")
        assert data.startswith("\ufeff")
        assert data.splitlines()[0].lstrip("\ufeff") == ";".join(CSV_HEADER)


class TestExcel:
    def _sheet(self):
        wb = load_workbook(io.BytesIO(to_xlsx(ITEMS, "Kadıköy · Otel", NOW)))
        return wb, wb["Kayıtlar"]

    def test_baslik_sabit_ve_filtreli(self):
        _, ws = self._sheet()
        assert [c.value for c in ws[1]] == CSV_HEADER
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref.startswith("A1:")

    def test_iletisim_hucreleri_tiklanabilir(self):
        _, ws = self._sheet()
        assert ws["C2"].hyperlink.target == "tel:+902163334455"
        assert ws["D2"].hyperlink.target == "mailto:info@ornek.com"
        assert ws["E2"].hyperlink.target == "https://www.ornek.com.tr/odalar?x=1"
        assert ws["H2"].value == "Haritada aç"
        assert ws["H2"].hyperlink.target.startswith("https://www.google.com/maps")

    def test_turkce_karakterler_korunur(self):
        _, ws = self._sheet()
        assert ws["A2"].value == "Şişli Öğretmen Evi 1"

    def test_bilgi_sayfasinda_baslik_ve_atif(self):
        wb, _ = self._sheet()
        info = wb["Bilgi"]
        assert info["A1"].value == "Kadıköy · Otel"
        assert "OpenStreetMap" in info["A5"].value


class TestPdf:
    def test_gecerli_pdf_ve_gomulu_yazi_tipi(self):
        """Standart Helvetica'da ş/ğ/ı yok; Inter gomulu olmali."""
        data = to_pdf(ITEMS, "Kadıköy · Karışık", NOW)
        assert data.startswith(b"%PDF")
        assert b"Inter" in data
        assert b"SpaceGrotesk" in data

    def test_bos_liste_de_pdf_uretir(self):
        assert to_pdf([], "Boş", NOW).startswith(b"%PDF")

    def test_tek_turde_tur_sutunu_yok(self):
        keys = [k for k, _, _ in pdf_columns(build_rows([_item(1), _item(2)]))]
        assert "type" not in keys

    def test_karisik_turde_tur_sutunu_var(self):
        keys = [k for k, _, _ in pdf_columns(build_rows(ITEMS))]
        assert "type" in keys

    def test_sutunlar_sayfaya_sigiyor(self):
        """Yatay A4 - kenar boslugu = 773.9 pt."""
        for rows in (build_rows(ITEMS), build_rows([_item(1)])):
            assert sum(w for _, _, w in pdf_columns(rows)) <= 773


class TestUc:
    """Uc: bicim alani, icerik tipi ve dosya uzantisi."""

    def _post(self, client, fmt=None):
        body = {"type": "hotel", "radius": 0, "center": {"lat": 0, "lon": 0}, "items": ITEMS}
        if fmt:
            body["format"] = fmt
        return client.post("/api/export", json=body)

    def test_varsayilan_csv(self, client):
        response = self._post(client)
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/csv")
        assert response.headers["content-disposition"].endswith('.csv"')

    def test_excel(self, client):
        response = self._post(client, "xlsx")
        assert response.status_code == 200
        assert "spreadsheetml" in response.headers["content-type"]
        assert response.headers["content-disposition"].endswith('.xlsx"')
        assert load_workbook(io.BytesIO(response.content))["Kayıtlar"]["A2"].value

    def test_pdf(self, client):
        response = self._post(client, "pdf")
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/pdf"
        assert response.content.startswith(b"%PDF")

    def test_bilinmeyen_bicim_reddedilir(self, client):
        assert self._post(client, "docx").status_code == 422

    def test_ust_sinir_10000(self, client):
        """Next route'u ile ayni sinir; asan istek dosya uretmeden reddedilir."""
        body = {
            "type": "hotel", "radius": 0, "center": {"lat": 0, "lon": 0},
            "items": [ITEMS[1]] * 10_001, "format": "csv",
        }
        assert client.post("/api/export", json=body).status_code == 422


class TestPdfSayfalama:
    """Elle sayfalama: hic satir kaybolmamali, sira bozulmamali."""

    def test_tum_satirlar_sirayla_tablolarda(self, monkeypatch):
        from reportlab.platypus import Table

        seen = []
        original = Table.__init__

        def spy(self, data, *args, **kwargs):
            seen.append(data)
            original(self, data, *args, **kwargs)

        monkeypatch.setattr(Table, "__init__", spy)
        to_pdf([_item(i) for i in range(1, 501)], "Sayfalama", NOW)

        numbers = [int(row[0]) for table in seen for row in table[1:]]
        assert numbers == list(range(1, 501))
        assert len(seen) > 1, "500 satir tek sayfaya sigmamali"
